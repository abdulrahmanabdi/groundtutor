"""Load PDF, PowerPoint, notes, and transcripts into chunks."""

from __future__ import annotations

import hashlib
import json
import re
from dataclasses import dataclass, field
from pathlib import Path

from groundtutor.chunking import Cue, chunk_cues, chunk_text, normalize_newlines
from groundtutor.errors import ConfigurationError, IngestError
from groundtutor.types import Chunk

_TS = re.compile(
    r"(?:(?P<h1>\d+):)?(?P<m1>\d{1,2}):(?P<s1>\d{2})[.,](?P<ms1>\d{1,3})\s*-->\s*"
    r"(?:(?P<h2>\d+):)?(?P<m2>\d{1,2}):(?P<s2>\d{2})[.,](?P<ms2>\d{1,3})"
)
_TAG = re.compile(r"<[^>]+>")
_VIDEO = (
    re.compile(r"(?:v=|vi=)([\w-]{11})"),
    re.compile(r"youtu\.be/([\w-]{11})"),
    re.compile(r"youtube\.com/(?:embed|shorts|live)/([\w-]{11})"),
    re.compile(r"^([\w-]{11})$"),
)
_DOC_ID = re.compile(r"[A-Za-z0-9][A-Za-z0-9._:-]{0,127}")

NOTES_SUFFIXES = {".txt", ".md", ".markdown"}
TRANSCRIPT_SUFFIXES = {".vtt", ".srt", ".json"}
PATH_SUFFIXES = {".pdf", ".pptx", *NOTES_SUFFIXES, *TRANSCRIPT_SUFFIXES}


@dataclass
class TextUnit:
    text: str
    page: int | None = None
    slide: int | None = None


@dataclass
class LoadedDocument:
    source_type: str
    source: str
    units: list[TextUnit] = field(default_factory=list)
    cues: list[Cue] = field(default_factory=list)


def make_doc_id(source: str, explicit: str | None = None, *, content: str = "") -> str:
    """Return a caller-supplied id, or a slug derived from the source name and content."""
    if explicit is not None:
        if not _DOC_ID.fullmatch(explicit):
            raise ValueError("doc_id must be 1-128 characters of letters, numbers, and . _ : -")
        return explicit
    label = source.split("\n", 1)[0]
    stem = Path(label).stem if label else "text"
    digest = hashlib.sha256(f"{source}\0{content}".encode()).hexdigest()[:8]
    slug = re.sub(r"[^A-Za-z0-9]+", "-", stem).strip("-").lower()[:40] or "doc"
    return f"{slug}-{digest}"


def load_path(path: Path) -> LoadedDocument:
    """Load a file based on its suffix."""
    path = path.expanduser()
    if not path.exists():
        raise IngestError(f"File not found: {path}")
    if not path.is_file():
        raise IngestError(f"Not a file: {path}")
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _load_pdf(path)
    if suffix == ".pptx":
        return _load_pptx(path)
    if suffix in NOTES_SUFFIXES:
        return _load_notes(path)
    if suffix in TRANSCRIPT_SUFFIXES:
        return load_transcript(path)
    supported = ", ".join(sorted(PATH_SUFFIXES))
    raise IngestError(f"Unsupported file type '{suffix}'. Expected one of: {supported}")


def load_transcript(
    source: str | Path,
    *,
    youtube_fetcher=None,
) -> LoadedDocument:
    """Load a VTT, SRT, or JSON transcript, or a YouTube URL or video id."""
    path = Path(source).expanduser()
    label = str(source)
    if path.exists():
        if not path.is_file():
            raise IngestError(f"Not a file: {path}")
        try:
            raw = path.read_text(encoding="utf-8-sig")
        except UnicodeDecodeError as exc:
            raise IngestError(f"{path} is not valid UTF-8") from exc
        cues = _cues_from_text(raw, path.suffix.lower())
        return LoadedDocument(source_type="transcript", source=str(path), cues=cues)
    video_id = parse_video_id(label)
    if video_id is None:
        raise IngestError(f"Transcript not found: {source}")
    fetcher = youtube_fetcher or fetch_youtube_cues
    cues = fetcher(video_id)
    if not isinstance(cues, list):
        cues = list(cues)
    if cues and not isinstance(cues[0], Cue):
        cues = cues_from_mappings(cues)
    if not cues:
        raise IngestError(f"No cues found for video {video_id}")
    return LoadedDocument(source_type="transcript", source=label, cues=list(cues))


def chunks_from_document(
    loaded: LoadedDocument,
    *,
    doc_id: str,
    chunk_size: int,
    overlap: int,
) -> list[Chunk]:
    """Chunk a loaded document, numbering pages and slides from 1."""
    if loaded.cues:
        chunks = chunk_cues(
            loaded.cues,
            doc_id=doc_id,
            source=loaded.source,
            chunk_size=chunk_size,
            overlap=overlap,
        )
    else:
        chunks = []
        for unit in loaded.units:
            chunks.extend(
                chunk_text(
                    unit.text,
                    doc_id=doc_id,
                    source=loaded.source,
                    source_type=loaded.source_type,
                    chunk_size=chunk_size,
                    overlap=overlap,
                    page=unit.page,
                    slide=unit.slide,
                    id_start=len(chunks),
                )
            )
    if not chunks:
        raise IngestError(f"No text found in {loaded.source}")
    return chunks


def parse_video_id(value: str) -> str | None:
    """Extract an 11-character YouTube id from a URL or a bare id."""
    text = value.strip()
    for pattern in _VIDEO:
        match = pattern.search(text)
        if match:
            return match.group(1)
    return None


def fetch_youtube_cues(video_id: str) -> list[Cue]:
    """Download a YouTube transcript. Requires the ``youtube`` extra."""
    try:
        from youtube_transcript_api import YouTubeTranscriptApi
    except ImportError as exc:
        raise ConfigurationError(
            'YouTube transcripts require the optional dependency: pip install "groundtutor[youtube]"'
        ) from exc
    api = YouTubeTranscriptApi()
    try:
        fetched = (
            api.fetch(video_id)
            if hasattr(api, "fetch")
            else YouTubeTranscriptApi.get_transcript(video_id)
        )
    except Exception as exc:
        raise IngestError(f"Could not fetch a transcript for {video_id}: {exc}") from exc
    return cues_from_mappings(fetched)


def cues_from_mappings(items) -> list[Cue]:
    """Convert ``{text, start, duration|end}`` records into cues."""
    cues: list[Cue] = []
    for item in items:
        if isinstance(item, Cue):
            cues.append(item)
            continue
        if isinstance(item, dict):
            text = item.get("text")
            start = item.get("start")
            end = item.get("end")
            duration = item.get("duration", 0.0)
        else:
            text = getattr(item, "text", None)
            start = getattr(item, "start", None)
            end = getattr(item, "end", None)
            duration = getattr(item, "duration", 0.0)
        if text is None or start is None:
            raise IngestError("Each cue needs 'text' and 'start'")
        start_s = float(start)
        end_s = float(end) if end is not None else start_s + float(duration or 0.0)
        piece = str(text).replace("\n", " ").strip()
        if piece:
            cues.append(Cue(piece, start_s, end_s))
    return cues


def _cues_from_text(raw: str, suffix: str) -> list[Cue]:
    stripped = raw.lstrip()
    if suffix == ".json" or stripped.startswith("[") or stripped.startswith("{"):
        return _cues_from_json(raw)
    return parse_transcript(raw)


def _cues_from_json(raw: str) -> list[Cue]:
    try:
        data = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise IngestError("Transcript JSON is invalid") from exc
    if isinstance(data, dict):
        data = data.get("cues") or data.get("transcript") or data.get("segments")
    if not isinstance(data, list):
        raise IngestError("Transcript JSON must be a list of cues")
    cues = cues_from_mappings(data)
    if not cues:
        raise IngestError("No cues found in transcript")
    return cues


def parse_transcript(text: str) -> list[Cue]:
    """Parse a WebVTT or SRT document into cues."""
    text = normalize_newlines(text.lstrip("\ufeff")).strip()
    cues: list[Cue] = []
    for block in re.split(r"\n\s*\n", text):
        lines = [line.strip() for line in block.split("\n") if line.strip()]
        match = None
        stamp_at = None
        for index, line in enumerate(lines):
            found = _TS.search(line)
            if found:
                match = found
                stamp_at = index
                break
        if match is None or stamp_at is None:
            continue
        payload = _unescape(_TAG.sub("", " ".join(lines[stamp_at + 1 :]))).strip()
        if not payload:
            continue
        cues.append(Cue(payload, _seconds(match, "1"), _seconds(match, "2")))
    if not cues:
        raise IngestError("No cues found in transcript")
    return cues


def _seconds(match: re.Match[str], side: str) -> float:
    hours = int(match.group(f"h{side}") or 0)
    minutes = int(match.group(f"m{side}"))
    seconds = int(match.group(f"s{side}"))
    millis = match.group(f"ms{side}")
    return hours * 3600 + minutes * 60 + seconds + int(millis) / (10 ** len(millis))


def _unescape(text: str) -> str:
    return (
        text.replace("&nbsp;", " ").replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")
    )


def _load_notes(path: Path) -> LoadedDocument:
    try:
        text = path.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError as exc:
        raise IngestError(f"{path} is not valid UTF-8") from exc
    kind = "markdown" if path.suffix.lower() in {".md", ".markdown"} else "text"
    return LoadedDocument(source_type=kind, source=str(path), units=[TextUnit(text=text)])


def _load_pdf(path: Path) -> LoadedDocument:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - declared dependency
        raise ConfigurationError("PDF support requires pypdf") from exc
    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            raise IngestError(f"{path} is encrypted")
        units = []
        for number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                units.append(TextUnit(text=text, page=number))
    except IngestError:
        raise
    except Exception as exc:
        raise IngestError(f"Could not read PDF {path}: {exc}") from exc
    if not units:
        raise IngestError(f"No extractable text in {path}")
    return LoadedDocument(source_type="pdf", source=str(path), units=units)


def _load_pptx(path: Path) -> LoadedDocument:
    try:
        from pptx import Presentation
    except ImportError as exc:  # pragma: no cover - declared dependency
        raise ConfigurationError("PowerPoint support requires python-pptx") from exc
    try:
        presentation = Presentation(str(path))
    except Exception as exc:
        raise IngestError(f"Could not read PowerPoint {path}: {exc}") from exc
    units: list[TextUnit] = []
    for number, slide in enumerate(presentation.slides, start=1):
        text = _slide_markdown(slide)
        if text:
            units.append(TextUnit(text=text, slide=number))
    if not units:
        raise IngestError(f"No extractable text in {path}")
    return LoadedDocument(source_type="pptx", source=str(path), units=units)


def _slide_markdown(slide) -> str:
    title = _slide_title(slide)
    body: list[str] = []
    for shape in slide.shapes:
        for line in _shape_lines(shape):
            if title and line == title and not body:
                continue
            body.append(line)
    parts: list[str] = []
    if title:
        parts.append(f"# {title}")
    if body:
        parts.append("\n".join(body))
    notes = _slide_notes(slide)
    if notes:
        parts.append(f"## Notes\n{notes}")
    return "\n\n".join(parts).strip()


def _slide_title(slide) -> str:
    try:
        title_shape = slide.shapes.title
    except (AttributeError, ValueError):
        return ""
    if title_shape is None:
        return ""
    return (title_shape.text or "").strip()


def _slide_notes(slide) -> str:
    if not slide.has_notes_slide:
        return ""
    frame = slide.notes_slide.notes_text_frame
    if frame is None:
        return ""
    return "\n".join(
        paragraph.text.strip() for paragraph in frame.paragraphs if paragraph.text.strip()
    )


def _shape_lines(shape) -> list[str]:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    lines: list[str] = []
    if getattr(shape, "has_text_frame", False):
        for paragraph in shape.text_frame.paragraphs:
            text = paragraph.text.strip()
            if text:
                lines.append(text)
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
        for child in shape.shapes:
            lines.extend(_shape_lines(child))
    return lines
