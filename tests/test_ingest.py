import json

import pytest
from pptx import Presentation
from pptx.util import Inches

from groundtutor.errors import IngestError
from groundtutor.ingest import (
    chunks_from_document,
    load_path,
    load_transcript,
    make_doc_id,
    parse_transcript,
    parse_video_id,
)
from tests.support import write_text_pdf

VTT = """WEBVTT

00:00:01.000 --> 00:00:04.000
Backpropagation <c>computes</c> gradients.

00:01.500 --> 00:04.000
The chain rule follows.
"""

SRT = """1
00:00:01,000 --> 00:00:04,000
Hello from the recording.

2
00:00:04,500 --> 00:00:08,000
Gradients flow backward.
"""


def test_parse_vtt_and_srt():
    vtt = parse_transcript(VTT)
    assert vtt[0].text == "Backpropagation computes gradients."
    assert vtt[0].start == pytest.approx(1.0)
    assert vtt[0].end == pytest.approx(4.0)
    assert vtt[1].start == pytest.approx(1.5)
    srt = parse_transcript(SRT)
    assert srt[0].text == "Hello from the recording."
    assert srt[1].start == pytest.approx(4.5)


def test_video_ids():
    assert parse_video_id("https://www.youtube.com/watch?v=dQw4w9WgXcQ") == "dQw4w9WgXcQ"
    assert parse_video_id("https://youtu.be/dQw4w9WgXcQ") == "dQw4w9WgXcQ"
    assert parse_video_id("https://www.youtube.com/shorts/dQw4w9WgXcQ") == "dQw4w9WgXcQ"
    assert parse_video_id("dQw4w9WgXcQ") == "dQw4w9WgXcQ"
    assert parse_video_id("not a video") is None


def test_pdf_pages_stay_separate(tmp_path):
    path = tmp_path / "lecture.pdf"
    write_text_pdf(
        path,
        [
            "Backpropagation computes gradients on page one.",
            "The chain rule is discussed on page two.",
        ],
    )
    loaded = load_path(path)
    chunks = chunks_from_document(loaded, doc_id="lecture", chunk_size=1000, overlap=50)
    pages = {chunk.page: chunk.text for chunk in chunks}
    assert 1 in pages and 2 in pages
    assert "page one" in pages[1]
    assert "page two" in pages[2]
    assert all(chunk.source_type == "pdf" for chunk in chunks)


def test_pptx_title_body_and_notes(tmp_path):
    path = tmp_path / "slides.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Chain rule"
    slide.placeholders[1].text = "Backpropagation computes gradients."
    slide.notes_slide.notes_text_frame.text = "Apply the chain rule at every layer."
    blank = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = blank.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    box.text_frame.text = "A weight matrix multiplies the input."
    presentation.save(path)

    loaded = load_path(path)
    chunks = chunks_from_document(loaded, doc_id="slides", chunk_size=1000, overlap=50)
    first = [chunk for chunk in chunks if chunk.slide == 1]
    second = [chunk for chunk in chunks if chunk.slide == 2]
    assert first and second
    assert any(chunk.heading_path[:1] == ("Chain rule",) for chunk in first)
    assert any("Backpropagation computes gradients." in chunk.text for chunk in first)
    assert any("chain rule" in chunk.text.lower() for chunk in first)
    assert any("weight matrix" in chunk.text for chunk in second)


def test_markdown_file_and_transcript_file(tmp_path):
    notes = tmp_path / "notes.md"
    notes.write_text("# Lesson\n\n## Gradients\n\nThey point uphill.\n", encoding="utf-8")
    loaded = load_path(notes)
    chunks = chunks_from_document(loaded, doc_id="notes", chunk_size=1000, overlap=20)
    assert chunks[-1].heading_path == ("Lesson", "Gradients")
    assert (
        notes.read_text(encoding="utf-8")[chunks[-1].char_start : chunks[-1].char_end]
        == chunks[-1].text
    )

    vtt = tmp_path / "talk.vtt"
    vtt.write_text(VTT, encoding="utf-8")
    transcript = load_path(vtt)
    timed = chunks_from_document(transcript, doc_id="talk", chunk_size=1000, overlap=20)
    assert timed[0].start_time == pytest.approx(1.0)
    assert "chain rule" in timed[0].text


def test_json_transcript_and_injected_youtube_fetcher(tmp_path):
    path = tmp_path / "talk.json"
    path.write_text(
        json.dumps([{"text": "Gradients flow backward.", "start": 12, "duration": 3}]),
        encoding="utf-8",
    )
    loaded = load_transcript(path)
    assert loaded.cues[0].start == 12
    assert loaded.cues[0].end == 15

    def fetcher(video_id: str):
        assert video_id == "dQw4w9WgXcQ"
        return [{"text": "A fetched line.", "start": 0, "duration": 1.5}]

    fetched = load_transcript("https://youtu.be/dQw4w9WgXcQ", youtube_fetcher=fetcher)
    assert fetched.cues[0].text == "A fetched line."
    assert fetched.cues[0].end == pytest.approx(1.5)


def test_missing_and_empty_documents(tmp_path):
    with pytest.raises(IngestError):
        load_path(tmp_path / "missing.pdf")
    empty = tmp_path / "empty.pdf"
    write_text_pdf(empty, [""])
    with pytest.raises(IngestError):
        load_path(empty)
    with pytest.raises(IngestError):
        load_transcript("this is not a video or a file")
    with pytest.raises(IngestError):
        parse_transcript("WEBVTT\n\nNOTE\njust a comment\n")


def test_doc_ids_change_with_content_and_reject_bad_explicit_ids():
    first = make_doc_id("inline.txt", content="alpha")
    second = make_doc_id("inline.txt", content="beta")
    assert first != second
    assert make_doc_id("lecture.pdf", "Lecture_1") == "Lecture_1"
    with pytest.raises(ValueError):
        make_doc_id("lecture.pdf", "has spaces")
